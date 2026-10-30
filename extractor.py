"""Document text extractor for PDF, Word (.docx), Excel (.xlsx), and plain text."""
import io
import os
import re


def extract_text_from_file(file_storage, filename: str) -> dict:
    """Extract text from an uploaded file object (Werkzeug FileStorage or file-like object).
    
    Returns a dict with:
        ok (bool)
        filename (str)
        suggested_title (str)
        text (str)
        word_count (int)
        file_type (str)
    """
    ext = os.path.splitext(filename)[1].lower()
    base_name = os.path.splitext(filename)[0]
    # Clean up filename for suggested topic title
    suggested_title = re.sub(r'[_\-\.]+', ' ', base_name).strip().title()

    file_bytes = file_storage.read()
    if not file_bytes:
        return {
            "ok": False,
            "error": "The uploaded file is empty."
        }

    buffer = io.BytesIO(file_bytes)
    extracted_text = ""
    file_type = ext.lstrip(".")

    try:
        if ext == ".pdf":
            import pypdf
            reader = pypdf.PdfReader(buffer)
            page_texts = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    page_texts.append(text.strip())
            extracted_text = "\n\n".join(page_texts)

        elif ext in [".docx", ".doc"]:
            import docx
            doc = docx.Document(buffer)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))
            
            extracted_text = "\n\n".join(paragraphs)

        elif ext in [".xlsx", ".xls", ".xlsm"]:
            import openpyxl
            wb = openpyxl.load_workbook(buffer, data_only=True)
            sheet_sections = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                sheet_lines = [f"=== Sheet: {sheetname} ==="]
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val).strip() for val in row if val is not None and str(val).strip()]
                    if row_vals:
                        sheet_lines.append(" | ".join(row_vals))
                if len(sheet_lines) > 1:
                    sheet_sections.append("\n".join(sheet_lines))
            extracted_text = "\n\n".join(sheet_sections)

        elif ext in [".txt", ".md", ".csv", ".json", ".tsv"]:
            # Try multiple encodings
            for enc in ["utf-8", "latin-1", "utf-16", "cp1252"]:
                try:
                    extracted_text = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if not extracted_text:
                extracted_text = file_bytes.decode("utf-8", errors="ignore")

        else:
            return {
                "ok": False,
                "error": f"Unsupported file format '{ext}'. Please upload a PDF, Word (.docx), Excel (.xlsx), or Text file."
            }

        extracted_text = extracted_text.strip()
        if not extracted_text:
            return {
                "ok": False,
                "error": "No readable text could be found in the document (it might be scanned images or empty)."
            }

        words = len(re.findall(r'\b\w+\b', extracted_text))
        return {
            "ok": True,
            "filename": filename,
            "suggested_title": suggested_title,
            "text": extracted_text,
            "word_count": words,
            "file_type": file_type
        }

    except Exception as e:
        return {
            "ok": False,
            "error": f"Failed to parse {filename}: {str(e)}"
        }

